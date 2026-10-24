"""Excel template, import (with preview) and export for products."""

from dataclasses import dataclass, field
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from apps.core.money import ZERO, q2, q3, to_decimal

from .models import Category, Product, QuickPreset, QuickSaleCategory, StockMovement
from .presets import PresetError, format_preset, parse_preset_list
from .stock import log_price_change, record_movement

COLUMNS = [
    "Name", "Barcode", "Category", "Sold by", "Selling price", "Cost price",
    "Opening stock", "Low-stock threshold", "Quick-sale category", "Preset buttons",
]

SOLD_BY_ALIASES = {
    "piece": Product.SoldBy.PIECE, "pieces": Product.SoldBy.PIECE, "pcs": Product.SoldBy.PIECE, "pc": Product.SoldBy.PIECE,
    "weight": Product.SoldBy.WEIGHT, "kg": Product.SoldBy.WEIGHT, "kgs": Product.SoldBy.WEIGHT,
    "volume": Product.SoldBy.VOLUME, "litre": Product.SoldBy.VOLUME, "liter": Product.SoldBy.VOLUME,
    "l": Product.SoldBy.VOLUME, "litres": Product.SoldBy.VOLUME,
}

EXAMPLE_ROWS = [
    ["Menengai Bar Soap 800g", "6161100110016", "Cleaning", "Piece", 210, 185, 24, 6, "", ""],
    ["Pishori Rice", "", "Rice", "Weight", 200, 170, 50, 10, "Rice", "1/4 kg; 1/2 kg; 1 kg; 2 kg"],
    ["Cooking Oil (loose)", "", "Cooking oil", "Volume", 300, 260, 20, 5, "Cooking oil", "KSh 20; KSh 50; KSh 100; 1/2 L; 1 L"],
]


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="11734B")
    for col, width in zip("ABCDEFGHIJ", [30, 16, 16, 10, 13, 12, 13, 18, 20, 36]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(COLUMNS)
    for row in EXAMPLE_ROWS:
        ws.append(row)
    dv = DataValidation(type="list", formula1='"Piece,Weight,Volume"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("D2:D5000")
    _style_header(ws)
    notes = wb.create_sheet("How to fill")
    for line in [
        ["Column", "What to put"],
        ["Name", "Required. Existing products are matched by barcode first, then by name."],
        ["Barcode", "Optional. Leave blank for loose goods."],
        ["Sold by", "Piece, Weight (priced per kg) or Volume (priced per litre)."],
        ["Selling price", "Required. Per piece, per kg or per litre."],
        ["Cost price", "Per piece, per kg or per litre."],
        ["Opening stock", "Only used for new products. Change existing stock with a recount or stock take."],
        ["Low-stock threshold", "Optional. Blank uses the shop default."],
        ["Quick-sale category", "Optional. Puts the product on the sell screen's quick-sale panel."],
        ["Preset buttons", "Optional, separated by ';'. Examples: 1/4 kg; 1/2 kg; 500 g; 1 L; KSh 50"],
    ]:
        notes.append(line)
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 90
    return _to_bytes(wb)


def export_products(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(COLUMNS + ["Current stock"])
    for p in queryset.select_related("category", "quick_category").prefetch_related("presets"):
        presets = "; ".join(format_preset(x.kind, x.value, p.unit) for x in p.presets.all() if not x.hidden)
        ws.append([
            p.name, p.barcode or "", p.category.name if p.category else "", p.get_sold_by_display().split(" ")[0],
            float(p.selling_price), float(p.cost_price), None,
            float(p.low_stock_threshold) if p.low_stock_threshold is not None else None,
            p.quick_category.name if p.quick_category else "", presets, float(p.stock_qty),
        ])
    _style_header(ws)
    return _to_bytes(wb)


def _to_bytes(wb):
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@dataclass
class ImportRow:
    row_number: int
    values: dict
    errors: list = field(default_factory=list)
    action: str = "create"
    product_id: int | None = None
    presets: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.errors

    def to_session(self):
        return {
            "row_number": self.row_number, "values": self.values, "errors": self.errors,
            "action": self.action, "product_id": self.product_id,
            "presets": [[k, str(v)] for k, v in self.presets],
        }


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_workbook(file):
    """Read an uploaded workbook and validate every row without saving anything."""
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for bad files
        raise ValueError(f"Could not read the file as an Excel workbook ({exc.__class__.__name__}).")
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header = [_cell(h).lower() for h in next(rows_iter, [])]
    index = {}
    for col in COLUMNS:
        if col.lower() in header:
            index[col] = header.index(col.lower())
    missing = [c for c in ("Name", "Sold by", "Selling price") if c not in index]
    if missing:
        raise ValueError(f"The first row must have the template's column names. Missing: {', '.join(missing)}.")

    existing_by_barcode = {p.barcode: p for p in Product.objects.exclude(barcode=None)}
    existing_by_name = {p.name.lower(): p for p in Product.objects.all()}
    seen_barcodes = {}
    seen_names = {}
    seen_products = {}
    results = []
    for number, raw in enumerate(rows_iter, start=2):
        values = {col: _cell(raw[i]) if i < len(raw) else "" for col, i in index.items()}
        if not any(values.values()):
            continue
        row = ImportRow(number, values)
        results.append(row)
        name = values.get("Name", "")
        barcode = values.get("Barcode", "")
        if not name:
            row.errors.append("Name is missing")
        sold_by = SOLD_BY_ALIASES.get(values.get("Sold by", "").lower())
        if sold_by is None:
            row.errors.append(f"Sold by must be Piece, Weight or Volume (got '{values.get('Sold by', '')}')")
        price = to_decimal(values.get("Selling price"))
        if price is None or price <= ZERO:
            row.errors.append("Selling price is missing or not a number")
        for col in ("Cost price", "Opening stock", "Low-stock threshold"):
            if values.get(col) and (to_decimal(values[col]) is None or to_decimal(values[col]) < ZERO):
                row.errors.append(f"{col} must be a number of zero or more")
        stock = to_decimal(values.get("Opening stock"))
        if sold_by == Product.SoldBy.PIECE and stock is not None and stock != stock.to_integral_value():
            row.errors.append("Opening stock must be a whole number for products sold by the piece")

        match = existing_by_barcode.get(barcode) if barcode else None
        by_name = existing_by_name.get(name.lower()) if name else None
        if match is None:
            match = by_name
        elif by_name is not None and by_name.pk != match.pk:
            row.errors.append(f"Barcode {barcode} belongs to '{match.name}', not '{by_name.name}'")
        if match is not None and barcode and barcode in existing_by_barcode and existing_by_barcode[barcode].pk != match.pk:
            row.errors.append(f"Barcode {barcode} already belongs to '{existing_by_barcode[barcode].name}'")
        if barcode:
            if barcode in seen_barcodes:
                row.errors.append(f"Barcode {barcode} is also on row {seen_barcodes[barcode]}")
            seen_barcodes.setdefault(barcode, number)
        if name:
            if name.lower() in seen_names:
                row.errors.append(f"'{name}' is also on row {seen_names[name.lower()]}")
            seen_names.setdefault(name.lower(), number)
        if match is not None:
            row.action, row.product_id = "update", match.pk
            if match.pk in seen_products:
                row.errors.append(f"Row {seen_products[match.pk]} already updates '{match.name}'")
            seen_products.setdefault(match.pk, number)

        if values.get("Preset buttons") and sold_by is not None:
            try:
                row.presets = parse_preset_list(values["Preset buttons"], Product.UNIT_LABELS[sold_by])
            except PresetError as exc:
                row.errors.append(f"Preset buttons: {exc}")
    return results


def rows_from_session(data):
    rows = []
    for item in data:
        row = ImportRow(item["row_number"], item["values"], item["errors"], item["action"], item["product_id"])
        row.presets = [(k, to_decimal(v)) for k, v in item["presets"]]
        rows.append(row)
    return rows


@transaction.atomic
def apply_rows(rows, user):
    """Create or update products from validated rows. Rows with errors are skipped."""
    created = updated = 0
    for row in rows:
        if not row.ok:
            continue
        v = row.values
        sold_by = SOLD_BY_ALIASES[v["Sold by"].lower()]
        product = Product.objects.get(pk=row.product_id) if row.product_id else Product()
        old_cost, old_price = product.cost_price or ZERO, product.selling_price or ZERO
        product.name = v["Name"]
        if v.get("Barcode"):
            product.barcode = v["Barcode"]
        product.sold_by = sold_by
        product.selling_price = q2(to_decimal(v["Selling price"]))
        if v.get("Cost price") and row.action == "create":
            product.cost_price = q2(to_decimal(v["Cost price"]))
        elif v.get("Cost price") and product.stock_qty <= ZERO:
            product.cost_price = q2(to_decimal(v["Cost price"]))
        if v.get("Category"):
            product.category, _ = Category.objects.get_or_create(name=v["Category"])
        if v.get("Low-stock threshold"):
            product.low_stock_threshold = q3(to_decimal(v["Low-stock threshold"]))
        if v.get("Quick-sale category"):
            product.quick_category, _ = QuickSaleCategory.objects.get_or_create(name=v["Quick-sale category"])
        product.save()
        log_price_change(product, old_cost=old_cost, old_price=old_price, user=user,
                         reason="Added from Excel" if row.action == "create" else "Updated from Excel")
        if row.presets:
            product.presets.all().delete()
            QuickPreset.objects.bulk_create(
                QuickPreset(product=product, kind=kind, value=value, order=i) for i, (kind, value) in enumerate(row.presets)
            )
        if row.action == "create":
            created += 1
            stock = to_decimal(v.get("Opening stock"))
            if stock and stock > ZERO:
                record_movement(
                    product, StockMovement.Kind.OPENING, stock, unit_cost=product.cost_price,
                    user=user, reason="Excel import",
                )
        else:
            updated += 1
    return created, updated
