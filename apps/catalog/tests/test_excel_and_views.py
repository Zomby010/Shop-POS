from decimal import Decimal
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.catalog import excel
from apps.catalog.models import Product, QuickPreset

HEADER = excel.COLUMNS


def workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile("products.xlsx", buffer.getvalue())


def test_template_has_columns_and_examples():
    wb = load_workbook(BytesIO(excel.build_template()))
    ws = wb["Products"]
    assert [c.value for c in ws[1]] == HEADER
    assert ws.max_row == 4


@pytest.mark.django_db
def test_parse_reports_row_errors(soap):
    rows = excel.parse_workbook(workbook([
        ["Pishori", "", "Rice", "Weight", 200, 170, 50, 10, "Rice", "1/4 kg; 1/2 kg; KSh 50"],
        ["", "", "", "Weight", 200, "", "", "", "", ""],
        ["Bad unit", "", "", "Box", 200, "", "", "", "", ""],
        ["No price", "", "", "Piece", "", "", "", "", "", ""],
        ["Other soap", "6161100110016", "", "Piece", 100, "", "", "", "", ""],
        ["Half bread", "", "", "Piece", 60, "", 2.5, "", "", ""],
        ["Oil", "", "", "Volume", 300, "", "", "", "Oil", "1 kg"],
        ["Bar soap", "", "", "Piece", 220, "", "", "", "", ""],
    ]))
    by_name = {r.values["Name"]: r for r in rows}
    assert by_name["Pishori"].ok and by_name["Pishori"].action == "create"
    assert len(by_name["Pishori"].presets) == 3
    assert "Name is missing" in by_name[""].errors
    assert any("Sold by" in e for e in by_name["Bad unit"].errors)
    assert any("Selling price" in e for e in by_name["No price"].errors)
    # Matched by barcode, so it renames the existing soap...
    assert by_name["Other soap"].ok and by_name["Other soap"].action == "update"
    assert any("whole number" in e for e in by_name["Half bread"].errors)
    assert any("Preset" in e for e in by_name["Oil"].errors)
    # ...and a second row aiming at the same product is flagged.
    assert any("already updates" in e for e in by_name["Bar soap"].errors)


@pytest.mark.django_db
def test_barcode_and_name_pointing_at_different_products(soap, make_product):
    make_product("Milk", barcode="222")
    rows = excel.parse_workbook(workbook([["Milk", "6161100110016", "", "Piece", 10, "", "", "", "", ""]]))
    assert any("belongs to" in e for e in rows[0].errors)


@pytest.mark.django_db
def test_duplicate_barcode_within_file():
    rows = excel.parse_workbook(workbook([
        ["A", "111", "", "Piece", 10, "", "", "", "", ""],
        ["B", "111", "", "Piece", 10, "", "", "", "", ""],
    ]))
    assert rows[0].ok and not rows[1].ok


@pytest.mark.django_db
def test_import_flow_preview_then_confirm(owner_client, soap):
    upload = workbook([
        ["Pishori", "", "Rice", "Weight", 200, 170, 50, 10, "Rice", "1/4 kg; 1/2 kg"],
        ["Bar soap", "6161100110016", "Cleaning", "Piece", 220, 190, 99, "", "", ""],
        ["Broken", "", "", "Box", 1, "", "", "", "", ""],
    ])
    response = owner_client.post("/products/import/", {"file": upload})
    assert response.status_code == 200
    assert b"2 ready" in response.content
    assert not Product.objects.filter(name="Pishori").exists()  # nothing saved on preview

    response = owner_client.post("/products/import/", {"confirm": "1"})
    assert response.status_code == 302
    rice = Product.objects.get(name="Pishori")
    assert rice.sold_by == Product.SoldBy.WEIGHT and rice.stock_qty == Decimal("50.000")
    assert rice.quick_category.name == "Rice"
    assert QuickPreset.objects.filter(product=rice).count() == 2
    soap.refresh_from_db()
    assert soap.selling_price == Decimal("220.00")
    assert soap.stock_qty == Decimal("24.000")  # existing stock is not touched by import


@pytest.mark.django_db
def test_export_downloads(owner_client, rice):
    response = owner_client.get("/products/export.xlsx")
    wb = load_workbook(BytesIO(response.content))
    assert wb.active["A2"].value == "Pishori"


@pytest.mark.django_db
def test_product_create_with_tiers_presets_and_opening_stock(owner_client):
    data = {
        "name": "Sugar", "barcode": "", "category": "", "sold_by": "weight", "selling_price": "180",
        "cost_price": "160", "low_stock_threshold": "", "quick_category": "", "quick_order": "0", "active": "on",
        "opening_stock": "50",
        "tiers-TOTAL_FORMS": "2", "tiers-INITIAL_FORMS": "0", "tiers-MIN_NUM_FORMS": "0", "tiers-MAX_NUM_FORMS": "1000",
        "tiers-0-min_qty": "10", "tiers-0-unit_price": "170",
        "presets-TOTAL_FORMS": "3", "presets-INITIAL_FORMS": "0", "presets-MIN_NUM_FORMS": "0", "presets-MAX_NUM_FORMS": "1000",
        "presets-0-text": "1/4 kg", "presets-0-order": "0",
        "presets-1-text": "KSh 50", "presets-1-order": "1",
    }
    response = owner_client.post("/products/new/", data)
    assert response.status_code == 302, response.context and response.context["form"].errors
    sugar = Product.objects.get(name="Sugar")
    assert sugar.stock_qty == Decimal("50.000")
    assert sugar.tiers.get().unit_price == Decimal("170")
    assert [(p.kind, p.value) for p in sugar.presets.all()] == [("qty", Decimal("0.250")), ("amount", Decimal("50.000"))]


@pytest.mark.django_db
def test_preset_with_wrong_unit_is_rejected(owner_client):
    data = {
        "name": "Oil", "sold_by": "volume", "selling_price": "300", "cost_price": "260", "quick_order": "0",
        "tiers-TOTAL_FORMS": "0", "tiers-INITIAL_FORMS": "0",
        "presets-TOTAL_FORMS": "1", "presets-INITIAL_FORMS": "0", "presets-0-text": "1 kg", "presets-0-order": "0",
    }
    response = owner_client.post("/products/new/", data)
    assert response.status_code == 200
    assert not Product.objects.filter(name="Oil").exists()


@pytest.mark.django_db
def test_cashier_cannot_edit_products(cashier_client, soap):
    assert cashier_client.get(f"/products/{soap.pk}/").status_code == 403
    assert cashier_client.get("/products/").status_code == 200
