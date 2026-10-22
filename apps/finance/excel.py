from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="11734B")


def workbook_response_bytes(title, header, rows, totals=None, subtitle=""):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    if subtitle:
        ws.append([subtitle])
    ws.append([])
    ws.append(header)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append([float(v) if hasattr(v, "as_tuple") else v for v in row])
    if totals:
        ws.append([float(v) if hasattr(v, "as_tuple") else v for v in totals])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, _ in enumerate(header, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = 18 if i > 1 else 28
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
