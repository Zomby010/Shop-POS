import zipfile
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.core import backup
from apps.core.models import ShopSettings
from apps.finance import reports
from apps.finance.models import Expense, ExpenseCategory
from apps.inventory import services as stock
from apps.partners.models import Customer
from apps.sales import services as sales


@pytest.fixture
def day_of_trading(cashier, owner, rice, soap, oil):
    """Three sales, one partial refund, one voided sale, an expense and a spillage."""
    def sell(items, **pay):
        cart = sales.current_cart(cashier)
        for product, kw in items:
            sales.add_to_cart(cart, product, **kw)
        return sales.complete_sale(cart, user=cashier, **pay)

    s1 = sell([(rice, {"quantity": Decimal("1")}), (soap, {"quantity": Decimal("2")})])  # 200 + 420 cash
    s2 = sell([(oil, {"amount": Decimal("100")})], mpesa="100")  # 100 M-Pesa
    customer = Customer.objects.create(name="Mama Wanjiku")
    sell([(soap, {"quantity": Decimal("1")})], credit="210", customer=customer)  # 210 deni
    s4 = sell([(rice, {"quantity": Decimal("2")})])  # 400, voided below
    sales.refund_lines(s1, {s1.lines.get(product=soap).pk: 1}, user=cashier, approver=owner, reason="Damaged")
    sales.void_sale(s4, user=cashier, approver=owner, reason="Mistake")
    Expense.objects.create(category=ExpenseCategory.objects.get_or_create(name="Transport")[0], amount=Decimal("150"), method="cash", user=owner)
    stock.record_loss(oil, "0.5", user=owner, reason="Spillage")
    return {"s1": s1, "s2": s2}


@pytest.mark.django_db
def test_profit_for_the_day(day_of_trading):
    today = timezone.localdate()
    total = reports.profit(today, today).total
    # Sales 620 + 100 + 210 + 400 = 1330; refunds 210 (soap) + 400 (void) = 610
    assert total.sales == Decimal("1330.00")
    assert total.refunds == Decimal("610.00")
    assert total.net_sales == Decimal("720.00")
    # Cost: rice 1kg 170 + soap 1 kept 185 + oil 0.333 L * 260 = 86.58 + soap 185 (deni) = 626.58
    assert total.cost_of_sales == Decimal("626.58")
    assert total.gross_profit == Decimal("93.42")
    assert total.expenses == Decimal("150.00")
    assert total.stock_loss == Decimal("130.00")  # 0.5 L oil at 260
    assert total.net_profit == Decimal("-186.58")


@pytest.mark.django_db
def test_sales_report_splits_payment_methods(day_of_trading):
    today = timezone.localdate()
    rows, total = reports.sales_report(today - timedelta(days=1), today)
    assert len(rows) == 2 and rows[0].count == 0
    assert total.count == 4 and total.voided == 1
    assert total.cash == Decimal("1020.00")  # 620 + 400 (later voided)
    assert total.mpesa == Decimal("100.00")
    assert total.credit == Decimal("210.00")
    assert total.refunds == Decimal("610.00")
    assert total.net == Decimal("720.00")


@pytest.mark.django_db
def test_best_and_slow_sellers(day_of_trading, make_product):
    unsold = make_product("Matches", price="10", cost="8", stock="100")
    today = timezone.localdate()
    best, slow, rows = reports.best_and_slow(today, today)
    by_name = {r.product.name: r for r in rows}
    assert by_name["Bar soap"].quantity == Decimal("2.000")  # 2 + 1 sold, 1 returned
    assert by_name["Bar soap"].revenue == Decimal("420.00")
    assert by_name["Pishori"].quantity == Decimal("1.000")  # voided sale excluded
    assert best[0].product.name == "Bar soap"
    assert slow[0].product == unsold


@pytest.mark.django_db
def test_stock_value(make_product):
    make_product("A", price="100", cost="80", stock="10")
    make_product("B", price="50", cost="40", stock="-2")  # negative stock is not counted as value
    rows, total = reports.stock_value()
    assert total.at_cost == Decimal("800.00") and total.at_price == Decimal("1000.00")
    assert total.potential_profit == Decimal("200.00")


@pytest.mark.django_db
@pytest.mark.parametrize("url", ["/money/", "/money/sales/", "/money/products/", "/money/stock-value/", "/money/expenses/"])
def test_report_pages_and_excel(owner_client, day_of_trading, url):
    assert owner_client.get(url).status_code == 200
    response = owner_client.get(url + "?format=xlsx")
    wb = load_workbook(BytesIO(response.content))
    assert wb.active.max_row >= 4


@pytest.mark.django_db
def test_reports_owner_only(cashier_client):
    assert cashier_client.get("/money/").status_code == 403


@pytest.mark.django_db
def test_add_and_delete_expense(owner_client):
    rent = ExpenseCategory.objects.get_or_create(name="Rent")[0]
    owner_client.post("/money/expenses/", {"date": timezone.localdate().isoformat(), "category": rent.pk,
                                          "amount": "15000", "method": "mpesa", "note": "October"})
    expense = Expense.objects.get()
    assert expense.amount == Decimal("15000.00")
    owner_client.post(f"/money/expenses/{expense.pk}/delete/")
    assert not Expense.objects.exists()


@pytest.mark.django_db(transaction=True)
def test_backup_zip_and_pruning(tmp_path, rice):
    shop = ShopSettings.load()
    shop.backup_folder = str(tmp_path / "usb")
    shop.backups_to_keep = 2
    shop.save()
    paths = []
    for i in range(3):
        paths.append(backup.create_backup())
        # make the next file name different
        paths[-1].rename(paths[-1].with_name(f"{backup.PREFIX}2026-01-0{i + 1}-200000.zip"))
    backup.create_backup()
    kept = backup.list_backups()
    assert len(kept) == 2
    with zipfile.ZipFile(kept[0]) as archive:
        assert "shoppos.sqlite3" in archive.namelist()


@pytest.mark.django_db
def test_backup_due(tmp_path, monkeypatch):
    shop = ShopSettings.load()
    shop.backup_folder = str(tmp_path)
    shop.save()
    tz = timezone.get_current_timezone()
    morning = timezone.make_aware(timezone.datetime(2026, 10, 22, 9, 0), tz)
    evening = timezone.make_aware(timezone.datetime(2026, 10, 22, 19, 0), tz)
    assert not backup.backup_due(morning)
    assert backup.backup_due(evening)
    (tmp_path / f"{backup.PREFIX}2026-10-22-183000.zip").write_bytes(b"x")
    assert not backup.backup_due(evening)


@pytest.mark.django_db(transaction=True)
def test_backup_page(owner_client, tmp_path):
    shop = ShopSettings.load()
    shop.backup_folder = str(tmp_path)
    shop.save()
    owner_client.post("/money/backup/")
    name = backup.list_backups()[0].name
    assert name.encode() in owner_client.get("/money/backup/").content
    response = owner_client.get(f"/money/backup/{name}")
    assert response.status_code == 200
    assert owner_client.get("/money/backup/..%2Fsecret_key.txt").status_code == 404


@pytest.mark.django_db
def test_setup_owner_command(monkeypatch, django_user_model):
    from django.core.management import call_command

    answers = iter(["shopowner"])
    secrets = iter(["secret12", "secret12", "4321"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    monkeypatch.setattr("getpass.getpass", lambda prompt="": next(secrets))
    call_command("setup_owner")
    owner = django_user_model.objects.get(username="shopowner")
    assert owner.check_password("secret12") and owner.profile.check_pin("4321")
    call_command("setup_owner")  # does nothing once a user exists
    assert django_user_model.objects.count() == 1
