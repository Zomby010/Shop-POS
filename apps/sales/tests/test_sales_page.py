from decimal import Decimal

import pytest
from openpyxl import load_workbook
from io import BytesIO

from apps.partners.models import Customer
from apps.sales import receipt, services
from apps.sales.models import Payment, Sale


@pytest.fixture
def three_sales(cashier, owner, soap, rice):
    def sell(product, qty, **pay):
        cart = services.current_cart(cashier)
        services.add_to_cart(cart, product, quantity=Decimal(qty))
        return services.complete_sale(cart, user=cashier, **pay)

    cash = sell(soap, "1")  # 210 cash
    mpesa = sell(rice, "2", mpesa="400", mpesa_code="sjk3h7qwe2")  # 400 M-Pesa
    customer = Customer.objects.create(name="Mama Wanjiku")
    deni = sell(soap, "2", mpesa="100", credit="320", customer=customer)  # split M-Pesa + deni
    services.refund_lines(cash, {cash.lines.get().pk: 1}, user=owner, approver=owner, reason="Damaged")
    return cash, mpesa, deni


@pytest.mark.django_db
def test_mpesa_code_is_saved_and_printed(three_sales):
    _, mpesa, _ = three_sales
    assert mpesa.payments.get().reference == "SJK3H7QWE2"
    text = "\n".join(r.left for r in receipt.receipt_rows(mpesa))
    assert "M-Pesa SJK3H7QWE2" in text


@pytest.mark.django_db
def test_pay_view_takes_mpesa_code(cashier_client, soap):
    cashier_client.post("/sell/cart/scan/", {"code": "6161100110016"})
    cashier_client.post("/sell/cart/pay/", {"mpesa": "210", "mpesa_code": "QWE123"}, HTTP_HX_REQUEST="true")
    assert Payment.objects.get().reference == "QWE123"


@pytest.mark.django_db
def test_sales_page_totals_and_filters(owner_client, three_sales):
    cash, mpesa, deni = three_sales
    response = owner_client.get("/sell/history/?period=today")
    totals = response.context["totals"]
    assert totals["count"] == 3 and totals["total"] == Decimal("1030.00")
    assert totals["cash"] == Decimal("210.00") and totals["mpesa"] == Decimal("500.00")
    assert totals["credit"] == Decimal("320.00") and totals["refunded"] == Decimal("210.00")

    def receipts(query):
        return {s.pk for s in owner_client.get(f"/sell/history/?{query}").context["page"]}

    assert receipts("method=cash") == {cash.pk}
    assert receipts("method=mpesa") == {mpesa.pk}
    assert receipts("method=credit") == {deni.pk}
    assert receipts("method=split") == {deni.pk}
    assert receipts("status=refunded") == {cash.pk}
    assert receipts("status=paid") == {mpesa.pk, deni.pk}
    assert receipts("q=qwe2") == {mpesa.pk}
    assert receipts("q=wanjiku") == {deni.pk}
    assert receipts("q=R000002&period=yesterday") == {mpesa.pk}  # receipt numbers ignore the dates
    assert receipts("period=yesterday") == set()


@pytest.mark.django_db
def test_sales_page_excel(owner_client, three_sales):
    response = owner_client.get("/sell/history/?format=xlsx")
    sheet = load_workbook(BytesIO(response.content)).active
    values = [c.value for row in sheet.iter_rows() for c in row]
    assert "SJK3H7QWE2" in values and "Part refunded" in values and "On deni" in values


@pytest.mark.django_db
def test_cashier_sees_only_own_sales(client, owner, three_sales, soap, django_user_model):
    from apps.core.roles import CASHIER
    from django.contrib.auth.models import Group

    other = django_user_model.objects.create_user("jane", password="x")
    other.groups.add(Group.objects.get_or_create(name=CASHIER)[0])
    client.force_login(other)
    assert client.get("/sell/history/").context["totals"]["count"] == 0
